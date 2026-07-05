# -*- coding: utf-8 -*-
"""ImageRights.uz — BMI himoyasi. Premium 12-slide deck builder (python-pptx)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.oxml.ns import qn
from copy import deepcopy

SP = os.path.dirname(os.path.abspath(__file__))
A  = os.path.join(SP, "assets")
def asset(n): return os.path.join(A, n)

# ---------------- palette ----------------
NAVY   = RGBColor(0x0B,0x12,0x20)
NAVY2  = RGBColor(0x0F,0x17,0x2A)
CARD_D = RGBColor(0x16,0x22,0x40)
CARD_D2= RGBColor(0x1B,0x29,0x4B)
STROKE_D=RGBColor(0x2B,0x3A,0x63)
BLUE   = RGBColor(0x33,0x66,0xFF)
BLUE4  = RGBColor(0x59,0x8D,0xFF)
BLUE2  = RGBColor(0xBC,0xD3,0xFF)
CYAN   = RGBColor(0x7C,0xC7,0xFF)
WHITE  = RGBColor(0xFF,0xFF,0xFF)
INK    = RGBColor(0x0F,0x17,0x2A)
INK2   = RGBColor(0x24,0x33,0x53)
MUTED_D= RGBColor(0x9D,0xB0,0xD8)
MUTED_D2=RGBColor(0x74,0x88,0xB4)
MUTED_L= RGBColor(0x5B,0x70,0x99)
CARD_L = WHITE
PANEL_L= RGBColor(0xF1,0xF4,0xFA)
STROKE_L=RGBColor(0xE1,0xE7,0xF3)
GREEN  = RGBColor(0x16,0xA3,0x4A)
GREEN_D= RGBColor(0x34,0xD3,0x99)
GREEN_BG=RGBColor(0xE7,0xF6,0xEC)
AMBER  = RGBColor(0xD9,0x88,0x0B)
RED    = RGBColor(0xE5,0x48,0x4D)
RED_D  = RGBColor(0xFF,0x6B,0x81)
RED_BG = RGBColor(0xFC,0xEA,0xEA)

# ---------------- fonts ----------------
F_BLACK="Montserrat Black"
F_XB   ="Montserrat ExtraBold"
F_HEAD ="Montserrat"          # use bold=True
F_SEMI ="Montserrat SemiBold"
F_KICK ="Space Grotesk Medium"
F_MONO ="Space Grotesk"       # bold=True -> Space Grotesk Bold
F_BODY ="Inter"
F_MED  ="Inter Medium"
F_SB   ="Inter SemiBold"

EMU_IN = 914400
def IN(v): return Inches(v)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5

# ---------------- low level helpers ----------------
def _alpha(color_elem, pct):
    a = color_elem.makeelement(qn('a:alpha'), {'val': str(int(pct*1000))})
    color_elem.append(a)

def set_fill(shape, color, alpha=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if alpha is not None:
        srgb = shape.fill.fore_color._xFill.find(qn('a:srgbClr'))
        _alpha(srgb, alpha)

def no_fill(shape): shape.fill.background()

def set_line(shape, color=None, w=1.0, alpha=None):
    if color is None:
        shape.line.fill.background(); return
    shape.line.color.rgb = color
    shape.line.width = Pt(w)
    if alpha is not None:
        srgb = shape.line.color._xFill.find(qn('a:srgbClr'))
        _alpha(srgb, alpha)

def no_shadow(shape):
    sp = shape._element.spPr
    if sp.find(qn('a:effectLst')) is None:
        sp.append(sp.makeelement(qn('a:effectLst'), {}))

def soft_shadow(shape, blur=0.14, dist=0.09, dir=5400000, alpha=76, color="0B1220"):
    sp = shape._element.spPr
    for e in sp.findall(qn('a:effectLst')): sp.remove(e)
    el = sp.makeelement(qn('a:effectLst'), {})
    sh = el.makeelement(qn('a:outerShdw'), {
        'blurRad': str(int(blur*EMU_IN)), 'dist': str(int(dist*EMU_IN)),
        'dir': str(dir), 'rotWithShape':'0'})
    c = sh.makeelement(qn('a:srgbClr'), {'val': color})
    c.append(c.makeelement(qn('a:alpha'), {'val': str(alpha*1000)}))
    sh.append(c); el.append(sh); sp.append(el)

def bg(slide, img):
    slide.shapes.add_picture(asset(img), 0, 0, prs.slide_width, prs.slide_height)

def rrect(slide,x,y,w,h,fill=None,line_c=None,line_w=1.0,rad=0.11,
          fill_alpha=None,line_alpha=None,shadow=False,shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, IN(x),IN(y),IN(w),IN(h))
    sp.shadow.inherit=False
    if fill is None: no_fill(sp)
    else: set_fill(sp, fill, fill_alpha)
    set_line(sp, line_c, line_w, line_alpha)
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            frac = max(0.02, min(0.5, rad/min(w,h)))
            sp.adjustments[0]=frac
        except Exception: pass
    if shadow: soft_shadow(sp)
    else: no_shadow(sp)
    sp.text_frame.paragraphs[0].text=""
    return sp

def oval(slide,x,y,w,h,fill=None,line_c=None,line_w=1.0,fill_alpha=None,line_alpha=None):
    sp=slide.shapes.add_shape(MSO_SHAPE.OVAL,IN(x),IN(y),IN(w),IN(h))
    sp.shadow.inherit=False
    if fill is None: no_fill(sp)
    else: set_fill(sp,fill,fill_alpha)
    set_line(sp,line_c,line_w,line_alpha); no_shadow(sp)
    return sp

def line(slide,x1,y1,x2,y2,color,w=1.0,alpha=None,dash=None):
    cn=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,IN(x1),IN(y1),IN(x2),IN(y2))
    cn.line.color.rgb=color; cn.line.width=Pt(w)
    if alpha is not None:
        srgb=cn.line.color._xFill.find(qn('a:srgbClr')); _alpha(srgb,alpha)
    if dash:
        d=cn.line._get_or_add_ln().makeelement(qn('a:prstDash'),{'val':dash})
        cn.line._get_or_add_ln().append(d)
    cn.shadow.inherit=False
    return cn

def _set_spc(run, spc):
    run.font._rPr.set('spc', str(int(spc*100)))

def textbox(slide,x,y,w,h,paras,anchor=MSO_ANCHOR.TOP,wrap=True):
    """paras: list of dict(runs=[dict(t,f,s,c,b,i,spc)], align, before, after, lh)"""
    tb=slide.shapes.add_textbox(IN(x),IN(y),IN(w),IN(h))
    tf=tb.text_frame; tf.word_wrap=wrap
    tf.vertical_anchor=anchor
    for m in ('left','right','top','bottom'):
        setattr(tf,'margin_'+m,0)
    for i,p in enumerate(paras):
        para = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        para.alignment=p.get('align',PP_ALIGN.LEFT)
        if p.get('before') is not None: para.space_before=Pt(p['before'])
        if p.get('after') is not None: para.space_after=Pt(p['after'])
        if p.get('lh') is not None: para.line_spacing=p['lh']
        for r in p['runs']:
            run=para.add_run(); run.text=r['t']
            run.font.name=r.get('f',F_BODY); run.font.size=Pt(r.get('s',18))
            run.font.bold=r.get('b',False); run.font.italic=r.get('i',False)
            run.font.color.rgb=r.get('c',INK)
            try: run.font.language_id=MSO_LANGUAGE_ID.UZBEK_LATIN
            except Exception: pass
            if r.get('spc'): _set_spc(run,r['spc'])
            # ensure latin+cs font
            rPr=run.font._rPr
            for tag in ('a:latin','a:cs'):
                e=rPr.find(qn(tag))
                if e is None:
                    e=rPr.makeelement(qn(tag),{}); rPr.append(e)
                e.set('typeface', r.get('f',F_BODY))
    return tb

def R(t,f=F_BODY,s=18,c=INK,b=False,i=False,spc=None):
    return dict(t=t,f=f,s=s,c=c,b=b,i=i,spc=spc)
def P(runs,align=PP_ALIGN.LEFT,before=None,after=None,lh=None):
    return dict(runs=runs,align=align,before=before,after=after,lh=lh)

def kicker(slide,x,y,text,color=BLUE4,tick=True,size=12.5):
    if tick:
        rrect(slide,x,y+0.02,0.20,0.20,fill=color,rad=0.05)
        tx=x+0.34
    else: tx=x
    textbox(slide,tx,y-0.06,6.0,0.4,[P([R(text.upper(),F_KICK,size,color,spc=3.2)])],
            anchor=MSO_ANCHOR.MIDDLE)

def pill(slide,x,y,w,h,text,fill,tcolor,size=11.5,font=F_KICK,line_c=None,spc=1.8,bold=False):
    sp=rrect(slide,x,y,w,h,fill=fill,line_c=line_c,line_w=1.0,rad=h/2)
    textbox(slide,x,y-0.02,w,h,[P([R(text,font,size,tcolor,b=bold,spc=spc)],align=PP_ALIGN.CENTER)],
            anchor=MSO_ANCHOR.MIDDLE)
    return sp

def footer(slide,idx,dark=True,total=12):
    c1 = MUTED_D2 if dark else MUTED_L
    c2 = MUTED_D if dark else MUTED_L
    line(slide,0.62,7.06,12.71,7.06, STROKE_D if dark else STROKE_L, 0.75, alpha=70 if dark else 100)
    textbox(slide,0.62,7.1,7,0.34,[P([
        R("ImageRights",F_SB,10.5,(WHITE if dark else INK)),
        R(".uz",F_SB,10.5,BLUE4),
        R("   ·   Jumabekov Samandar  ·  BMI himoyasi 2026",F_MED,10.5,c1)])],
        anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide,10.7,7.1,2.01,0.34,[P([
        R(f"{idx:02d}",F_MONO,11,(WHITE if dark else INK),b=True),
        R(f" / {total:02d}",F_MONO,11,c2)],align=PP_ALIGN.RIGHT)],anchor=MSO_ANCHOR.MIDDLE)

def brandmark(slide,x,y,scale=1.0,dark=True):
    s=0.5*scale
    tile=rrect(slide,x,y,s,s,fill=BLUE,rad=0.13*scale)
    # scan frame glyph
    m=s*0.24; ix,iy=x+m,y+m; iw=s-2*m
    tick=iw*0.32
    for (cx,cy,dx,dy) in [(ix,iy,1,1),(ix+iw,iy,-1,1),(ix,iy+iw,1,-1),(ix+iw,iy+iw,-1,-1)]:
        line(slide,cx,cy,cx+dx*tick,cy,WHITE,1.6)
        line(slide,cx,cy,cx,cy+dy*tick,WHITE,1.6)
    oval(slide,ix+iw/2-s*0.05,iy+iw/2-s*0.05,s*0.1,s*0.1,fill=WHITE)
    textbox(slide,x+s+0.14,y-0.03,3.4,s+0.06,[P([
        R("ImageRights",F_SB,15*scale,(WHITE if dark else INK)),
        R(".uz",F_SB,15*scale,BLUE4)])],anchor=MSO_ANCHOR.MIDDLE)

# ---- monoline glyphs inside a tile ----
def glyph(slide,kind,x,y,s,color):
    cx,cy=x+s/2,y+s/2
    lw=1.9
    if kind=="scan":
        t=s*0.26; o=s*0.14
        for (gx,gy,dx,dy) in [(x+o,y+o,1,1),(x+s-o,y+o,-1,1),(x+o,y+s-o,1,-1),(x+s-o,y+s-o,-1,-1)]:
            line(slide,gx,gy,gx+dx*t,gy,color,lw); line(slide,gx,gy,gx,gy+dy*t,color,lw)
        oval(slide,cx-s*0.07,cy-s*0.07,s*0.14,s*0.14,fill=color)
    elif kind=="check":
        oval(slide,x+s*0.12,y+s*0.12,s*0.76,s*0.76,line_c=color,line_w=lw)
        line(slide,cx-s*0.17,cy+s*0.01,cx-s*0.03,cy+s*0.15,color,lw+0.3)
        line(slide,cx-s*0.03,cy+s*0.15,cx+s*0.20,cy-s*0.14,color,lw+0.3)
    elif kind=="doc":
        rrect(slide,x+s*0.22,y+s*0.12,s*0.56,s*0.76,line_c=color,line_w=lw,rad=s*0.06)
        for k,yy in enumerate([0.34,0.5,0.66]):
            line(slide,x+s*0.33,y+s*yy,x+s*(0.67 if k<2 else 0.55),y+s*yy,color,lw)
    elif kind=="globe":
        oval(slide,x+s*0.14,y+s*0.14,s*0.72,s*0.72,line_c=color,line_w=lw)
        oval(slide,x+s*0.36,y+s*0.14,s*0.28,s*0.72,line_c=color,line_w=lw)
        line(slide,x+s*0.14,cy,x+s*0.86,cy,color,lw)
    elif kind=="lock":
        rrect(slide,x+s*0.24,cy-s*0.02,s*0.52,s*0.42,line_c=color,line_w=lw,rad=s*0.05)
        # shackle
        oval(slide,x+s*0.34,y+s*0.18,s*0.32,s*0.36,line_c=color,line_w=lw)
        rrect(slide,x+s*0.24,cy-s*0.02,s*0.52,s*0.30,fill=None,line_c=color,line_w=lw,rad=s*0.05)
    elif kind=="spark":
        line(slide,cx,y+s*0.16,cx,y+s*0.84,color,lw)
        line(slide,x+s*0.16,cy,x+s*0.84,cy,color,lw)
        line(slide,x+s*0.28,y+s*0.28,x+s*0.72,y+s*0.72,color,lw*0.7)
        line(slide,x+s*0.72,y+s*0.28,x+s*0.28,y+s*0.72,color,lw*0.7)
    elif kind=="alert":
        line(slide,cx,y+s*0.16,x+s*0.86,y+s*0.80,color,lw)
        line(slide,x+s*0.86,y+s*0.80,x+s*0.14,y+s*0.80,color,lw)
        line(slide,x+s*0.14,y+s*0.80,cx,y+s*0.16,color,lw)
        line(slide,cx,y+s*0.40,cx,y+s*0.60,color,lw+0.4)
        oval(slide,cx-s*0.02,y+s*0.66,s*0.05,s*0.05,fill=color)
    elif kind=="eye":
        oval(slide,x+s*0.12,y+s*0.30,s*0.76,s*0.40,line_c=color,line_w=lw)
        oval(slide,cx-s*0.11,cy-s*0.11,s*0.22,s*0.22,fill=color)
    elif kind=="scale":  # licensing / commerce (use bag-ish)
        rrect(slide,x+s*0.24,y+s*0.30,s*0.52,s*0.52,line_c=color,line_w=lw,rad=s*0.05)
        oval(slide,x+s*0.36,y+s*0.16,s*0.28,s*0.28,line_c=color,line_w=lw)
    elif kind=="gavel":  # legal doc-check
        rrect(slide,x+s*0.22,y+s*0.14,s*0.50,s*0.72,line_c=color,line_w=lw,rad=s*0.05)
        line(slide,x+s*0.31,cy+s*0.06,x+s*0.40,cy+s*0.16,color,lw)
        line(slide,x+s*0.40,cy+s*0.16,x+s*0.58,cy-s*0.10,color,lw)
    elif kind=="fingerprint":
        for i,rr in enumerate([0.34,0.24,0.14]):
            oval(slide,cx-s*rr,cy-s*rr,s*2*rr,s*2*rr,line_c=color,line_w=lw)

def icon_tile(slide,x,y,s,kind,tile_fill,gcolor,rad=None,line_c=None):
    rrect(slide,x,y,s,s,fill=tile_fill,rad=(rad if rad else s*0.28),line_c=line_c)
    glyph(slide,kind,x+s*0.04,y+s*0.04,s*0.92,gcolor)

def chevron(slide,cx,cy,size,color,alpha=100):
    line(slide,cx-size*0.35,cy-size*0.5,cx+size*0.35,cy,color,2.2,alpha=alpha)
    line(slide,cx+size*0.35,cy,cx-size*0.35,cy+size*0.5,color,2.2,alpha=alpha)

def picture(slide,img,x,y,w=None,h=None):
    kw={}
    if w: kw['width']=IN(w)
    if h: kw['height']=IN(h)
    return slide.shapes.add_picture(asset(img),IN(x),IN(y),**kw)

# ============================================================
# SLIDE 1 — COVER
# ============================================================
def slide1():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_dark_C.png")
    # face mesh hero, right side
    picture(s,"facemesh_hero.png",8.05,0.25,w=5.0)
    brandmark(s,0.62,0.55,scale=1.0,dark=True)
    # top-right chip
    pill(s,9.9,0.66,2.85,0.4,"FACIAL DATA · CONSENT",fill=None,tcolor=CYAN,size=10.5,line_c=STROKE_D,spc=2.2)
    # kicker
    textbox(s,0.64,1.9,7.3,0.4,[P([R("TOSHKENT DAVLAT YURIDIK UNIVERSITETI  ·  FUQAROLIK HUQUQI SHO‘BASI",F_KICK,11.5,BLUE4,spc=2.4)])])
    # title
    textbox(s,0.6,2.38,7.9,2.6,[
        P([R("Tasvirga bo‘lgan huquqni",F_XB,32,WHITE)],lh=1.06,after=2),
        P([R("fuqarolik-huquqiy tartibga solish",F_XB,32,WHITE)],lh=1.06,after=2),
        P([R("va ",F_XB,32,WHITE),R("himoya qilish",F_XB,32,BLUE4),R(" masalalari",F_XB,32,WHITE)],lh=1.06),
    ])
    # divider
    line(s,0.64,4.9,3.0,4.9,BLUE,2.4)
    textbox(s,0.64,5.12,7.4,0.5,[P([R("Bitiruv malakaviy ishi himoyasi",F_SB,16.5,BLUE2)])])
    # author block
    textbox(s,0.64,5.74,7.6,1.0,[
        P([R("MUALLIF   ",F_KICK,10.5,MUTED_D2,spc=2.0),R("Jumabekov Samandar Bahromjon o‘g‘li",F_SB,13.5,WHITE)],after=5),
        P([R("RAHBAR   ",F_KICK,10.5,MUTED_D2,spc=2.0),R("yuridik fanlar bo‘yicha PhD  B. U. Koryog‘diyev",F_SB,13.5,WHITE)]),
    ])
    # three-result mini chips bottom
    labs=[("Ilmiy natija",BLUE4),("Normativ natija",CYAN),("Amaliy natija — ImageRights.uz",GREEN_D)]
    cxp=0.64
    for t,c in labs:
        w=0.28+len(t)*0.083
        sp=rrect(s,cxp,6.5,w,0.4,fill=CARD_D,rad=0.2,line_c=STROKE_D,fill_alpha=60)
        oval(s,cxp+0.16,6.66,0.09,0.09,fill=c)
        textbox(s,cxp+0.34,6.48,w,0.4,[P([R(t,F_MED,11,WHITE)])],anchor=MSO_ANCHOR.MIDDLE)
        cxp+=w+0.2
    # HUD labels near mesh
    for tx,ty,tt in [(11.45,2.15,"IMAGE RIGHTS"),(8.85,5.32,"ID · 0xA7F2"),(11.5,4.6,"MATCH 0.98")]:
        textbox(s,tx,ty,1.9,0.3,[P([R(tt,F_MONO,10,CYAN,spc=1.5)])])
    footer(s,1,dark=True)

# ============================================================
# SLIDE 2 — CENTRAL PROBLEM
# ============================================================
def slide2():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_dark_R.png")
    kicker(s,0.64,0.62,"Markaziy muammo")
    textbox(s,0.6,1.12,5.3,2.6,[
        P([R("O‘zbekiston qonunchiligi fuqaroga o‘z tasviri ustidan ",F_XB,23,WHITE),
           R("samarali nazorat",F_XB,23,BLUE4),
           R(" o‘rnatish imkonini beradimi?",F_XB,23,WHITE)],lh=1.12),
    ])
    textbox(s,0.64,3.62,4.95,1.1,[P([R("FK 99-modda tasvirga bo‘lgan huquqni tan oladi — biroq amaliyotda to‘rt tugun uni himoyasiz qoldiradi:",F_MED,13,MUTED_D)],lh=1.32)])
    # mesh accent bottom-left
    picture(s,"facemesh_mini.png",0.55,4.55,w=2.15)
    # 4 issue cards, right 2x2
    items=[
        ("01","scan","Huquqiy ta’rif yo‘q","Tasvirning maxsus qonuniy ta’rifi FKda mavjud emas."),
        ("02","check","Rozilik chegarasi","Rozilik shakllari va uning amal qilish doirasi belgilanmagan."),
        ("03","spark","AI va deepfake","Sun’iy intellekt keltirgan yangi, oldin bo‘lmagan xavflar."),
        ("04","alert","Tezkor himoya yo‘q","Internetdagi huquqbuzarlikda shoshilinch chora yetishmaydi."),
    ]
    x0,y0,cw,ch,gx,gy=6.05,1.15,3.28,2.62,0.34,0.3
    for i,(n,gl,t,d) in enumerate(items):
        cx=x0+(i%2)*(cw+gx); cy=y0+(i//2)*(ch+gy)
        rrect(s,cx,cy,cw,ch,fill=CARD_D,rad=0.16,line_c=STROKE_D,fill_alpha=88)
        icon_tile(s,cx+0.28,cy+0.28,0.62,gl,CARD_D2,BLUE4,line_c=STROKE_D)
        textbox(s,cx+1.05,cy+0.28,cw-1.2,0.62,[P([R(n,F_MONO,20,BLUE4,b=True)])],anchor=MSO_ANCHOR.MIDDLE)
        textbox(s,cx+0.3,cy+1.12,cw-0.6,0.5,[P([R(t,F_HEAD,15.5,WHITE,b=True)])])
        textbox(s,cx+0.3,cy+1.62,cw-0.58,0.9,[P([R(d,F_BODY,11.5,MUTED_D)],lh=1.28)])
    footer(s,2,dark=True)

# ============================================================
# SLIDE 3 — MAIN LEGAL GAP
# ============================================================
def slide3():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_dark_R.png")
    kicker(s,0.64,0.62,"Asosiy huquqiy bo‘shliq")
    textbox(s,0.6,1.1,12,0.7,[P([R("Huquq bor — mexanizm yo‘q",F_XB,27,WHITE)])])
    # left: FK 99 recognized
    lx,ly,lw,lh=0.64,2.2,4.35,3.05
    rrect(s,lx,ly,lw,lh,fill=CARD_D,rad=0.18,line_c=STROKE_D,fill_alpha=90)
    pill(s,lx+0.34,ly+0.34,1.9,0.42,"FK 99-MODDA",fill=None,tcolor=BLUE4,size=11.5,line_c=STROKE_D,spc=2)
    icon_tile(s,lx+lw-1.02,ly+0.28,0.62,"check",RGBColor(0x11,0x2E,0x24),GREEN_D)
    textbox(s,lx+0.34,ly+1.15,lw-0.6,1.0,[
        P([R("Huquq ",F_XB,23,WHITE),R("tan olingan",F_XB,23,GREEN_D)],after=4),
        P([R("Tasvirga bo‘lgan huquq shaxsiy nomulkiy huquqlar qatorida sanab o‘tilgan.",F_MED,13,MUTED_D)],lh=1.3)])
    pill(s,lx+0.34,ly+lh-0.72,2.55,0.42,"NOMODDIY NE’MAT SIFATIDA",fill=RGBColor(0x11,0x2E,0x24),tcolor=GREEN_D,size=10,spc=1.5)
    # center BUT
    textbox(s,5.05,2.55,1.95,1.2,[P([R("LEKIN",F_BLACK,30,WHITE)],align=PP_ALIGN.CENTER)],anchor=MSO_ANCHOR.MIDDLE)
    chevron(s,6.02,4.05,0.34,BLUE4)
    # right: missing mechanism
    rx,ry,rw,rh=7.05,2.2,5.65,3.05
    rrect(s,rx,ry,rw,rh,fill=CARD_D,rad=0.18,line_c=RED,fill_alpha=90,line_alpha=55)
    icon_tile(s,rx+0.34,ry+0.3,0.62,"alert",RGBColor(0x2E,0x14,0x18),RED_D)
    textbox(s,rx+1.12,ry+0.3,rw-1.3,0.62,[P([R("Amalga oshirish mexanizmi yetarli emas",F_HEAD,15.5,WHITE,b=True)],lh=1.05)],anchor=MSO_ANCHOR.MIDDLE)
    miss=["Tasvir ta’rifi","Rozilik tartibi","Istisno holatlar","Tijorat maqsadida foydalanish","Internetdagi himoya"]
    mx,my=rx+0.34,ry+1.2
    ww=(rw-0.68-0.24)/2
    for i,m in enumerate(miss):
        col=i%2; row=i//2
        px=mx+col*(ww+0.24); py=my+row*0.56
        if i==4: px=mx; ww2=rw-0.68
        else: ww2=ww
        rrect(s,px,py,ww2,0.46,fill=RGBColor(0x24,0x12,0x18),rad=0.1,line_c=RED,line_alpha=45)
        line(s,px+0.2,py+0.16,px+0.34,py+0.30,RED_D,2); line(s,px+0.34,py+0.16,px+0.2,py+0.30,RED_D,2)
        textbox(s,px+0.5,py,ww2-0.6,0.46,[P([R(m,F_SB,11.5,RGBColor(0xF0,0xC8,0xCC))])],anchor=MSO_ANCHOR.MIDDLE)
    # bottom tagline
    rrect(s,0.64,5.55,12.06,0.92,fill=CARD_D,rad=0.16,line_c=STROKE_D,fill_alpha=70)
    textbox(s,1.0,5.55,11.35,0.92,[P([
        R("“Qonunchilik lakunasi”",F_XB,16,BLUE4),
        R("  —  tasvirning mazmuni, chegaralari, rozilik tartibi va istisnolarini tartibga soluvchi alohida norma mavjud emas.",F_MED,13.5,WHITE)],lh=1.25)],anchor=MSO_ANCHOR.MIDDLE)
    footer(s,3,dark=True)

# ============================================================
# SLIDE 4 — LEGAL NATURE
# ============================================================
def slide4():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_light_main.png")
    kicker(s,0.64,0.62,"Huquqiy tabiat",color=BLUE)
    textbox(s,0.6,1.12,12.1,1.2,[
        P([R("Tasvir — shaxsni individuallashtiruvchi ",F_XB,25,INK),R("nomoddiy ne’mat",F_XB,25,BLUE)],lh=1.06,after=2),
        P([R("Tasvirga bo‘lgan huquq — shaxsiy nomulkiy huquq",F_SB,16,MUTED_L)],lh=1.1),
    ])
    # concept diagram: SHAXSIY MANFAAT <-> TASVIR <-> IQTISODIY QIYMAT
    dy=2.95; dh=1.9
    def node(x,w,title,sub,fill,tc,strong=False):
        rrect(s,x,dy,w,dh,fill=fill,rad=0.16,line_c=(BLUE if strong else STROKE_L),line_w=(1.6 if strong else 1.1),shadow=True)
        textbox(s,x+0.2,dy+0.34,w-0.4,0.9,[P([R(title,F_HEAD,15.5,tc,b=True)],align=PP_ALIGN.CENTER,lh=1.05)],anchor=MSO_ANCHOR.MIDDLE)
        textbox(s,x+0.2,dy+dh-0.72,w-0.4,0.6,[P([R(sub,F_BODY,10.5,(BLUE2 if strong else MUTED_L))],align=PP_ALIGN.CENTER,lh=1.15)],anchor=MSO_ANCHOR.MIDDLE)
    node(0.64,3.55,"SHAXSIY MANFAAT","Sha’n · qadr-qimmat · daxlsizlik",CARD_L,INK)
    node(4.9,3.55,"TASVIR",  "identifikatsiyalovchi nomoddiy ne’mat",NAVY2,WHITE,strong=True)
    node(9.15,3.55,"IQTISODIY QIYMAT","litsenziya · reklama · homiylik",CARD_L,INK)
    for cxp in (4.55,8.8):
        chevron(s,cxp,dy+dh/2,0.3,BLUE)
    # dual tags above outer nodes
    pill(s,1.29,2.52,2.25,0.36,"extra-patrimonial · axloqiy",fill=RGBColor(0xEC,0xF1,0xFF),tcolor=BLUE,size=9,spc=0.6)
    pill(s,9.8,2.52,2.25,0.36,"patrimonial · mulkiy",fill=RGBColor(0xEC,0xF1,0xFF),tcolor=BLUE,size=9,spc=0.6)
    pill(s,5.5,2.52,2.35,0.36,"droit à l’image",fill=NAVY2,tcolor=BLUE2,size=9,spc=0.6)
    # conclusion card
    rrect(s,0.64,5.28,12.06,1.2,fill=NAVY2,rad=0.16,shadow=True)
    icon_tile(s,1.0,5.58,0.62,"lock",CARD_D2,BLUE4)
    textbox(s,1.95,5.28,10.5,1.2,[P([
        R("Huquqning o‘zi begonalashtirilmaydi",F_XB,15.5,WHITE),
        R("  —  ammo tasvirdan foydalanish vakolati shartnoma asosida berilishi mumkin (dualistik model).",F_MED,13.5,BLUE2)],lh=1.3)],anchor=MSO_ANCHOR.MIDDLE)
    footer(s,4,dark=False)

# ============================================================
# SLIDE 5 — COMPARATIVE LAW
# ============================================================
def slide5():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_light_main.png")
    kicker(s,0.64,0.62,"Qiyosiy-huquqiy tajriba",color=BLUE)
    textbox(s,0.6,1.12,12.1,0.7,[P([R("Xorijiy tajribadan ",F_XB,25,INK),R("huquqiy saboqlar",F_XB,25,BLUE)])])
    cards=[
        ("FR","Fransiya",["droit à l’image — mutlaq huquq","référé: bir necha soatda tezkor sud himoyasi"]),
        ("DE","Germaniya",["KUG 22–23: rozilik — umumiy qoida","aniq belgilangan istisnolar"]),
        ("PL","Polsha",["FK 23-modda: shaxsiy nomulkiy manfaat","ikki qonunda parallel himoya"]),
        ("RU","Rossiya",["FK 152¹: maxsus tartibga solish","internetda olib tashlash talabi"]),
        ("EU","GDPR",["tasvir — identifikatsiyalovchi ma’lumot","unutilish huquqi (17-modda)"]),
    ]
    x0=0.64; y0=2.1; cw=2.3; ch=3.0; g=(12.06-5*cw)/4
    for i,(code,name,pts) in enumerate(cards):
        cx=x0+i*(cw+g)
        rrect(s,cx,y0,cw,ch,fill=CARD_L,rad=0.15,line_c=STROKE_L,shadow=True)
        rrect(s,cx,y0,cw,0.12,fill=BLUE,rad=0.06)
        rrect(s,cx+0.26,y0+0.34,0.7,0.52,fill=NAVY2,rad=0.1)
        textbox(s,cx+0.26,y0+0.32,0.7,0.52,[P([R(code,F_MONO,15,WHITE,b=True,spc=0.5)],align=PP_ALIGN.CENTER)],anchor=MSO_ANCHOR.MIDDLE)
        textbox(s,cx+1.02,y0+0.32,cw-1.1,0.56,[P([R(name,F_HEAD,15,INK,b=True)])],anchor=MSO_ANCHOR.MIDDLE)
        line(s,cx+0.26,y0+1.06,cx+cw-0.26,y0+1.06,STROKE_L,1)
        yy=y0+1.24
        for p in pts:
            oval(s,cx+0.28,yy+0.07,0.1,0.1,fill=BLUE4)
            textbox(s,cx+0.5,yy-0.04,cw-0.72,0.85,[P([R(p,F_MED,10.5,INK2)],lh=1.2)])
            yy+=0.86
    # conclusion strip
    rrect(s,0.64,5.42,12.06,1.06,fill=NAVY2,rad=0.16,shadow=True)
    pill(s,1.0,5.72,2.15,0.46,"O‘ZBEKISTON UCHUN",fill=BLUE,tcolor=WHITE,size=10.5,spc=1.2,bold=True)
    textbox(s,3.4,5.42,9.1,1.06,[P([R("Kontinental model — tasvirni ",F_MED,13.5,BLUE2),R("shaxsiy nomulkiy huquq",F_SB,13.5,WHITE),R(" sifatida maxsus norma bilan mustahkamlash eng maqbul yondashuv.",F_MED,13.5,BLUE2)],lh=1.28)],anchor=MSO_ANCHOR.MIDDLE)
    footer(s,5,dark=False)

# ============================================================
# SLIDE 6 — DEEPFAKE
# ============================================================
def slide6():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_dark_C.png")
    kicker(s,0.64,0.62,"Yangi raqamli tahdid",color=RED_D)
    textbox(s,0.6,1.08,8.2,1.2,[
        P([R("Muammo endi faqat ruxsatsiz",F_XB,25,WHITE)],lh=1.04,after=2),
        P([R("surat tarqatish ",F_XB,25,WHITE),R("emas",F_XB,25,RED_D)],lh=1.04)])
    # morph mesh center-right
    picture(s,"facemesh_deepfake.png",7.5,1.4,w=5.3)
    textbox(s,8.3,1.95,1.8,0.3,[P([R("AUTHENTIC",F_MONO,10,BLUE4,spc=1.5)],align=PP_ALIGN.CENTER)])
    textbox(s,10.15,1.95,2.0,0.3,[P([R("MANIPULATED",F_MONO,10,RED_D,spc=1.5)],align=PP_ALIGN.CENTER)])
    line(s,9.85,1.9,9.85,2.35,MUTED_D2,0.9,alpha=55,dash="dash")
    # sequence chips (left)
    seq=[("AI",BLUE4),("DEEPFAKE",BLUE4),("SOXTA KONTEKST",AMBER),("REPUTATSION VA MULKIY ZARAR",RED_D)]
    yy=2.55
    for i,(t,c) in enumerate(seq):
        w=0.4+len(t)*0.092
        rrect(s,0.64,yy,min(w,4.9),0.5,fill=CARD_D,rad=0.1,line_c=STROKE_D,fill_alpha=90)
        oval(s,0.84,yy+0.19,0.12,0.12,fill=c)
        textbox(s,1.1,yy,4.6,0.5,[P([R(t,F_SB,12.5,WHITE)])],anchor=MSO_ANCHOR.MIDDLE)
        if i<3: chevron(s,0.95,yy+0.62,0.2,MUTED_D2)
        yy+=0.62
    # thesis stat chips
    textbox(s,5.25,2.62,2.2,1.6,[
        P([R("96%",F_MONO,26,RED_D,b=True)],after=0),
        P([R("deepfake video —\npornografik (Sensity AI)",F_BODY,10.5,MUTED_D)],lh=1.2,after=8),
        P([R("+550%",F_MONO,22,AMBER,b=True)],after=0),
        P([R("deepfake hajmi\n2019–2024",F_BODY,10.5,MUTED_D)],lh=1.2)])
    # legal questions bottom
    qs=[("gavel","Kim javobgar?"),("globe","Platformaning majburiyati qanday?"),("alert","Kontent qancha tez olib tashlanishi kerak?")]
    x0=0.64; y0=5.35; cw=(12.06-2*0.3)/3
    for i,(gl,q) in enumerate(qs):
        cx=x0+i*(cw+0.3)
        rrect(s,cx,y0,cw,1.15,fill=CARD_D,rad=0.14,line_c=STROKE_D,fill_alpha=92)
        icon_tile(s,cx+0.26,y0+0.28,0.58,gl,CARD_D2,CYAN)
        textbox(s,cx+1.0,y0+0.2,cw-1.2,0.78,[P([R(q,F_HEAD,13.5,WHITE,b=True)],lh=1.08)],anchor=MSO_ANCHOR.MIDDLE)
    footer(s,6,dark=True)

# ============================================================
# SLIDE 7 — MAIN PROPOSAL
# ============================================================
def slide7():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_light_main.png")
    kicker(s,0.64,0.62,"Mening asosiy taklifim",color=BLUE)
    textbox(s,0.6,1.08,8.6,1.4,[
        P([R("Fuqarolik kodeksiga ",F_XB,23,INK),R("maxsus norma",F_XB,23,BLUE)],lh=1.05,after=2),
        P([R("“Fuqaro tasviridan foydalanishni huquqiy tartibga solish”",F_SB,14.5,MUTED_L)],lh=1.12)])
    # original contribution badge
    pill(s,9.7,1.15,3.0,0.5,"ORIGINAL ILMIY HISSA",fill=NAVY2,tcolor=WHITE,size=11,spc=1.8,bold=True)
    textbox(s,9.7,1.72,3.0,0.4,[P([R("FK 19, 99-modda + yangi 178-modda",F_MED,10.5,MUTED_L)],align=PP_ALIGN.CENTER)])
    comps=[
        ("01","doc","Tasvirning huquqiy ta’rifi","Identifikatsiyalovchi vizual axborot sifatida."),
        ("02","check","Rozilik shakllari","Yozma · og‘zaki · elektron rozilik."),
        ("03","scan","Roziliksiz foydalanish istisnolari","Davlat, jamoat va umumiy plan hollari."),
        ("04","scale","Tijorat maqsadida foydalanish","Faqat yozma rozilik asosida."),
        ("05","eye","Videokuzatuv qoidalari","Ogohlantirish belgisi · daxlsiz zonalar."),
        ("06","lock","Maxsus himoya vositalari","Olib tashlash · shoshilinch chora · zarar."),
    ]
    x0=0.64; y0=2.72; cw=(12.06-2*0.34)/3; ch=1.78; gx=0.34; gy=0.28
    for i,(n,gl,t,d) in enumerate(comps):
        cx=x0+(i%3)*(cw+gx); cy=y0+(i//3)*(ch+gy)
        rrect(s,cx,cy,cw,ch,fill=CARD_L,rad=0.15,line_c=STROKE_L,shadow=True)
        rrect(s,cx,cy,0.14,ch,fill=BLUE,rad=0.07)
        icon_tile(s,cx+0.32,cy+0.3,0.6,gl,PANEL_L,BLUE)
        textbox(s,cx+cw-1.0,cy+0.28,0.85,0.5,[P([R(n,F_MONO,18,BLUE2,b=True)],align=PP_ALIGN.RIGHT)])
        textbox(s,cx+0.34,cy+1.0,cw-0.6,0.44,[P([R(t,F_HEAD,13.5,INK,b=True)],lh=1.02)])
        textbox(s,cx+0.34,cy+1.42,cw-0.55,0.34,[P([R(d,F_BODY,10,MUTED_L)],lh=1.15)])
    footer(s,7,dark=False)

# ============================================================
# SLIDE 8 — CONSENT & EXCEPTIONS
# ============================================================
def slide8():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_light_main.png")
    kicker(s,0.64,0.62,"Rozilik modeli",color=BLUE)
    textbox(s,0.6,1.08,12,0.7,[P([R("Rozilik — asosiy qoida, istisnolar — aniq belgilangan",F_XB,22,INK)])])
    # left: general rule + forms
    lx,ly,lw=0.64,2.0,5.75
    rrect(s,lx,ly,lw,2.15,fill=CARD_L,rad=0.16,line_c=STROKE_L,shadow=True)
    pill(s,lx+0.32,ly+0.3,1.5,0.4,"UMUMIY QOIDA",fill=GREEN_BG,tcolor=GREEN,size=10,spc=1.5)
    textbox(s,lx+0.32,ly+0.86,lw-0.6,0.8,[P([R("Fuqaro tasviridan foydalanish uchun ",F_SB,14.5,INK),R("rozilik talab qilinadi",F_XB,14.5,GREEN)],lh=1.2)])
    forms=["Yozma","Og‘zaki","Elektron"]
    fw=(lw-0.64-0.3)/3
    for i,f in enumerate(forms):
        fx=lx+0.32+i*(fw+0.15)
        rrect(s,fx,ly+1.6,fw,0.42,fill=PANEL_L,rad=0.1,line_c=STROKE_L)
        textbox(s,fx,ly+1.6,fw,0.42,[P([R(f,F_SB,11.5,INK2)],align=PP_ALIGN.CENTER)],anchor=MSO_ANCHOR.MIDDLE)
    # right: exceptions
    rx=6.62; rw=6.08
    rrect(s,rx,ly,rw,2.15,fill=CARD_L,rad=0.16,line_c=STROKE_L,shadow=True)
    pill(s,rx+0.32,ly+0.3,3.15,0.4,"ROZILIKSIZ — ISTISNO HOLATLAR",fill=RGBColor(0xEC,0xF1,0xFF),tcolor=BLUE,size=9.5,spc=1.0)
    exs=["Davlat va jamoat manfaatlari","Xavfsizlik va qidiruv","Ommaviy kasbiy faoliyat","Umumiy plan (asosiy obyekt emas)"]
    ew=(rw-0.64-0.24)/2
    for i,e in enumerate(exs):
        ex=rx+0.32+(i%2)*(ew+0.24); ey=ly+0.9+(i//2)*0.56
        rrect(s,ex,ey,ew,0.46,fill=PANEL_L,rad=0.1)
        oval(s,ex+0.18,ey+0.15,0.16,0.16,line_c=BLUE,line_w=1.6)
        textbox(s,ex+0.46,ey,ew-0.56,0.46,[P([R(e,F_MED,10.5,INK2)],lh=1.0)],anchor=MSO_ANCHOR.MIDDLE)
    # critical rule banner
    rrect(s,0.64,4.42,12.06,0.86,fill=NAVY2,rad=0.14,shadow=True)
    icon_tile(s,1.0,4.58,0.54,"scale",CARD_D2,AMBER)
    textbox(s,1.75,4.42,10.8,0.86,[P([R("TIJORAT MAQSADIDA FOYDALANISH",F_XB,15,WHITE,spc=0.5),R("   →   ",F_XB,15,AMBER),R("YOZMA ROZILIK SHART",F_XB,15,GREEN_D,spc=0.5)],lh=1.1)],anchor=MSO_ANCHOR.MIDDLE)
    # memorable conclusion emphasis
    rrect(s,0.64,5.5,12.06,0.98,fill=CARD_L,rad=0.14,line_c=BLUE,line_w=1.5,shadow=True)
    textbox(s,0.9,5.5,11.6,0.98,[P([
        R("Suratga tushishga rozilik",F_XB,18,INK),
        R("   ≠   ",F_BLACK,20,RED),
        R("cheklanmagan tarqatishga rozilik",F_XB,18,BLUE)],align=PP_ALIGN.CENTER,lh=1.05)],anchor=MSO_ANCHOR.MIDDLE)
    footer(s,8,dark=False)

# ============================================================
# SLIDE 9 — URGENT DIGITAL PROTECTION
# ============================================================
def slide9():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_dark_L.png")
    kicker(s,0.64,0.62,"Tezkor raqamli himoya",color=CYAN)
    textbox(s,0.6,1.08,12,1.0,[P([R("Raqamli muhitda ",F_XB,25,WHITE),R("kechikkan himoya",F_XB,25,RED_D),R(" — samarasiz himoya",F_XB,25,WHITE)],lh=1.05)])
    textbox(s,0.64,1.95,11.8,0.5,[P([R("Kontent soniyalarda tarqaladi — oddiy, sekin protsedura yetarli emas. Zanjir bir butun bo‘lishi kerak:",F_MED,13,MUTED_D)],lh=1.25)])
    # process timeline (5 steps)
    steps=[("alert","HUQUQBUZARLIK"),("doc","ELEKTRON DALILNI SAQLASH"),("gavel","SHOSHILINCH SUD CHORASI"),("scan","TARQALISHNI TO‘XTATISH"),("check","ZARARNI QOPLASH")]
    n=len(steps); x0=0.64; y0=2.75; tw=12.06; cw=(tw-(n-1)*0.32)/n; ch=1.55
    for i,(gl,t) in enumerate(steps):
        cx=x0+i*(cw+0.32)
        strong = i in (1,2)
        rrect(s,cx,y0,cw,ch,fill=(CARD_D2 if strong else CARD_D),rad=0.14,line_c=(CYAN if strong else STROKE_D),line_alpha=(60 if strong else 100),fill_alpha=95)
        icon_tile(s,cx+cw/2-0.31,y0+0.24,0.62,gl,NAVY2,(CYAN if strong else BLUE4))
        textbox(s,cx+0.12,y0+0.96,cw-0.24,0.5,[P([R(t,F_SB,10.5,WHITE)],align=PP_ALIGN.CENTER,lh=1.05)],anchor=MSO_ANCHOR.MIDDLE)
        if i<n-1: chevron(s,cx+cw+0.16,y0+ch/2,0.22,CYAN)
        # step number
        textbox(s,cx+0.12,y0+0.16,cw-0.24,0.3,[P([R(f"0{i+1}",F_MONO,9.5,MUTED_D2,spc=1)],align=PP_ALIGN.CENTER)])
    # référé reference card
    rrect(s,0.64,4.65,7.4,1.7,fill=CARD_D,rad=0.16,line_c=STROKE_D,fill_alpha=92)
    pill(s,0.98,4.95,1.55,0.4,"FRANSIYA · RÉFÉRÉ",fill=NAVY2,tcolor=CYAN,size=10,spc=1.2,line_c=STROKE_D)
    textbox(s,0.98,5.5,6.8,0.8,[P([R("Sud ",F_SB,13.5,WHITE),R("bir necha soat ichida",F_XB,13.5,CYAN),R(" tasvirni olib tashlash choralarini qo‘llaydi. FPKga shunga o‘xshash shoshilinch tartib joriy etish taklif etiladi.",F_MED,12.5,MUTED_D)],lh=1.28)])
    # virality propagation visual
    rrect(s,8.24,4.65,4.46,1.7,fill=CARD_D,rad=0.16,line_c=STROKE_D,fill_alpha=92)
    textbox(s,8.5,4.8,4,0.34,[P([R("VIRALLIK — TARQALISH",F_KICK,10,MUTED_D2,spc=1.8)])])
    ox,oy=8.75,5.75
    oval(s,ox-0.13,oy-0.13,0.26,0.26,fill=RED_D)
    import math
    for k in range(9):
        ang=k*0.7+0.2; dist=0.7+ (k%3)*0.42
        ex=ox+dist*1.25; ey=oy-0.55+ (k*0.11)
        ex=min(ex,12.45)
        line(s,ox,oy,ex,ey,BLUE4,1.0,alpha=45)
        oval(s,ex-0.05,ey-0.05,0.1,0.1,fill=BLUE4,fill_alpha=90)
    footer(s,9,dark=True)

# ============================================================
# SLIDE 10 — TRANSITION TO SOLUTION
# ============================================================
def slide10():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_dark_C.png")
    kicker(s,0.64,0.62,"Tadqiqotdan amaliy yechimga",color=GREEN_D)
    # formula
    fx=0.64; fy=1.2
    parts=[("ILMIY MUAMMO",BLUE4),("HUQUQIY TAKLIF",CYAN),("TEXNOLOGIK YECHIM",GREEN_D)]
    xx=fx
    for i,(t,c) in enumerate(parts):
        w=0.5+len(t)*0.108
        rrect(s,xx,fy,w,0.56,fill=CARD_D,rad=0.12,line_c=STROKE_D,fill_alpha=90)
        textbox(s,xx,fy,w,0.56,[P([R(t,F_SB,12.5,c,spc=0.5)],align=PP_ALIGN.CENTER)],anchor=MSO_ANCHOR.MIDDLE)
        xx+=w
        if i<2:
            chevron(s,xx+0.16,fy+0.28,0.24,MUTED_D); xx+=0.44
    # big reveal
    textbox(s,0.6,2.35,7.7,1.5,[
        P([R("ImageRights",F_BLACK,52,WHITE),R(".uz",F_BLACK,52,BLUE4)],lh=1.0)])
    textbox(s,0.64,3.5,7.3,1.0,[P([R("Tasvir huquqlarini himoya qilishga qaratilgan ",F_MED,15,BLUE2),R("LegalTech platforma prototipi",F_SB,15,WHITE),R(".",F_MED,15,BLUE2)],lh=1.3)])
    # tagline emphasis
    rrect(s,0.64,4.65,7.3,1.0,fill=CARD_D,rad=0.16,line_c=STROKE_D,fill_alpha=80)
    textbox(s,0.95,4.65,6.8,1.0,[P([R("Tadqiqot ",F_SB,15,MUTED_D),R("tavsiyalar bilan tugamadi",F_XB,15,GREEN_D),R(" — u ishlaydigan prototipga aylandi.",F_SB,15,MUTED_D)],lh=1.28)],anchor=MSO_ANCHOR.MIDDLE)
    # result badges
    for i,(t,c) in enumerate([("Ilmiy",BLUE4),("Normativ",CYAN),("Amaliy",GREEN_D)]):
        bx=0.64+i*1.5
        pill(s,bx,6.0,1.35,0.44,t,fill=NAVY2,tcolor=c,size=11,spc=1,line_c=STROKE_D)
    # dashboard mockup right
    picture(s,"mock_dashboard.png",7.75,1.85,w=5.55)
    footer(s,10,dark=True)

# ============================================================
# SLIDE 11 — HOW IT WORKS
# ============================================================
def slide11():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_light_main.png")
    kicker(s,0.64,0.56,"Platforma qanday ishlaydi",color=BLUE)
    textbox(s,0.6,1.02,12.1,0.6,[P([R("Tadqiqot takliflari ",F_XB,22,INK),R("platforma modullariga",F_XB,22,BLUE),R(" bog‘landi",F_XB,22,INK)])])
    # flow 6 steps
    flow=["TASVIRNI QAYD ETISH","ELEKTRON ROZILIK","MONITORING VA AI TAHLIL","ELEKTRON DALIL","HUQUQIY HIMOYA","LITSENZIYALASH"]
    n=len(flow); x0=0.64; y0=1.78; tw=12.06; cw=(tw-(n-1)*0.26)/n
    for i,t in enumerate(flow):
        cx=x0+i*(cw+0.26)
        rrect(s,cx,y0,cw,0.72,fill=(NAVY2 if i in(1,2,3) else CARD_L),rad=0.11,line_c=(STROKE_L),shadow=True)
        textbox(s,cx+0.12,y0,cw-0.24,0.72,[P([R(f"{i+1}",F_MONO,11,(BLUE4 if i in(1,2,3) else BLUE),b=True)],align=PP_ALIGN.CENTER,after=1),
            P([R(t,F_SB,8.6,(WHITE if i in(1,2,3) else INK2))],align=PP_ALIGN.CENTER,lh=0.95)],anchor=MSO_ANCHOR.MIDDLE)
        if i<n-1: chevron(s,cx+cw+0.13,y0+0.36,0.16,BLUE4)
    # left: mapping thesis -> module
    mx,my=0.64,2.86; mw=6.15
    rrect(s,mx,my,mw,3.05,fill=CARD_L,rad=0.16,line_c=STROKE_L,shadow=True)
    textbox(s,mx+0.32,my+0.24,mw-0.6,0.4,[P([R("TAKLIF  →  MODUL",F_KICK,11,MUTED_L,spc=2)])])
    maps=[("Elektron rozilik taklifi","Consent management moduli","check"),
          ("Deepfake muammosi","AI tahlil konsepti","spark"),
          ("Tezkor himoya zarurati","Dalil + huquqiy hujjat oqimi","doc"),
          ("Tasvirning iqtisodiy qiymati","Litsenziyalash modeli","scale")]
    yy=my+0.78
    for a,b,gl in maps:
        icon_tile(s,mx+0.32,yy,0.5,gl,PANEL_L,BLUE)
        textbox(s,mx+0.96,yy-0.03,mw-1.3,0.56,[
            P([R(a,F_SB,11.5,INK)],after=1),
            P([R(b,F_MED,10.5,BLUE)],lh=1.0)])
        yy+=0.56
    # right: two screenshots
    picture(s,"mock_consent.png",6.98,2.68,w=5.72)
    picture(s,"mock_deepfake.png",9.62,4.42,w=3.35)
    # disclaimer
    rrect(s,0.64,6.02,6.15,0.5,fill=RGBColor(0xFF,0xF7,0xE8),rad=0.1,line_c=RGBColor(0xF0,0xDD,0xB0))
    textbox(s,0.86,6.02,5.8,0.5,[P([R("⚠ Prototip sud, advokat yoki davlat organi o‘rnini bosmaydi — himoya jarayonini tezlashtiradi.",F_MED,9.2,RGBColor(0x8A,0x63,0x12))],lh=1.05)],anchor=MSO_ANCHOR.MIDDLE)
    footer(s,11,dark=False)

# ============================================================
# SLIDE 12 — THREE RESULTS / CONCLUSION
# ============================================================
def slide12():
    s=prs.slides.add_slide(BLANK)
    bg(s,"bg_dark_C.png")
    kicker(s,0.64,0.62,"Uch natija",color=GREEN_D)
    textbox(s,0.6,1.08,12,0.7,[P([R("Bir tadqiqot — ",F_XB,25,WHITE),R("uch bog‘langan natija",F_XB,25,BLUE4)])])
    pillars=[("01","ILMIY NATIJA","scan",BLUE4,"Tasvirga bo‘lgan huquqning huquqiy tabiati nomoddiy ne’mat sifatida asoslandi."),
             ("02","NORMATIV NATIJA","doc",CYAN,"Fuqarolik kodeksi va huquqni qo‘llash amaliyotiga aniq takliflar ishlab chiqildi."),
             ("03","AMALIY NATIJA","spark",GREEN_D,"ImageRights.uz LegalTech platforma prototipi yaratildi.")]
    x0=0.64; y0=2.1; cw=(12.06-2*0.36)/3; ch=2.7
    for i,(n,t,gl,c,d) in enumerate(pillars):
        cx=x0+i*(cw+0.36)
        rrect(s,cx,y0,cw,ch,fill=CARD_D,rad=0.18,line_c=STROKE_D,fill_alpha=94,shadow=False)
        rrect(s,cx,y0,cw,0.13,fill=c,rad=0.06)
        textbox(s,cx+0.34,y0+0.34,cw-0.6,0.5,[P([R(n,F_MONO,20,c,b=True)])])
        icon_tile(s,cx+cw-1.02,y0+0.32,0.64,gl,NAVY2,c)
        textbox(s,cx+0.34,y0+1.08,cw-0.6,0.5,[P([R(t,F_HEAD,16.5,WHITE,b=True,spc=0.3)])])
        line(s,cx+0.34,y0+1.62,cx+cw-0.34,y0+1.62,STROKE_D,1)
        textbox(s,cx+0.34,y0+1.76,cw-0.64,0.9,[P([R(d,F_MED,12,MUTED_D)],lh=1.3)])
    # final statement
    rrect(s,0.64,5.12,12.06,1.28,fill=NAVY2,rad=0.18,line_c=BLUE,line_w=1.4,shadow=False)
    textbox(s,0.9,5.12,11.55,1.28,[P([
        R("ANIQ HUQUQIY NORMA",F_XB,15,WHITE),
        R("  +  ",F_BLACK,15,BLUE4),
        R("TEZKOR HIMOYA MEXANIZMI",F_XB,15,CYAN),
        R("  +  ",F_BLACK,15,BLUE4),
        R("ZAMONAVIY TEXNOLOGIK YECHIM",F_XB,15,GREEN_D)],align=PP_ALIGN.CENTER,lh=1.25)],anchor=MSO_ANCHOR.MIDDLE)
    footer(s,12,dark=True)

for f in [slide1,slide2,slide3,slide4,slide5,slide6,slide7,slide8,slide9,slide10,slide11,slide12]:
    f()

out=os.path.join(SP,"ImageRights_BMI_Final.pptx")
prs.save(out)
print("SAVED", out)
